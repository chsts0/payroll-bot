"""Расчёт и разбор входных данных (таблица и ручной ввод)."""
import csv
import io
import math
import re
from collections import OrderedDict
from dataclasses import dataclass, field

RATE = 320               # ставка по умолчанию (для ручного ввода)
KPI_PERCENT = 4          # КПИ = 4% от продаж
ROUND_TO = 50            # «К выплате» округляется вверх до 50

# код в таблице -> (подпись на карточке, фон плашки, цвет текста)
POINTS = {
    "ДЗ": ("Завод", "#e3f5e9", "#1f7a44"),
    "3В": ("3 вокз", "#dff3f3", "#16706f"),
    "ЛД": ("Лесн", "#fde6ea", "#b3264a"),
    "С": ("Севк", "#efe3fa", "#7a2fb8"),
    "В": ("Вокз", "#e6effc", "#2b5aa8"),
    "П": ("Пик", "#fdeee3", "#b8520e"),
}

_ALIASES = {
    "ДЗ": ["дз", "дизайнзавод", "завод"],
    "3В": ["3в", "тривокзала", "3вокзала", "депотривокзала", "депо3вокзала", "3вокз"],
    "ЛД": ["лд", "лесная", "депонлесной", "деполесная", "лесн"],
    "С": ["с", "севкабель", "севк"],
    "В": ["в", "вокзал", "вокз"],
    "П": ["п", "пик"],
}


def _norm(s: str) -> str:
    s = str(s).lower().replace("ё", "е")
    s = s.translate(str.maketrans({"з": "3", "c": "с", "b": "в", "p": "р", "n": "п"})) if len(s.strip()) <= 2 else s
    return re.sub(r"[\s.\-_«»\"']", "", s)


ALIAS_MAP = {}
for _code, _al in _ALIASES.items():
    for _a in _al + [_code]:
        ALIAS_MAP[_norm(_a)] = _code


def find_point(s):
    return ALIAS_MAP.get(_norm(s))


def norm_name(s: str) -> str:
    return " ".join(str(s).replace("ё", "е").replace("Ё", "Е").split()).lower()


@dataclass
class Entry:
    point: str
    hours: float
    sales: float
    rate: float = RATE
    option: float = 0
    ofzp: float = 0
    expected: float = None      # «к выплате» из таблицы, для сверки

    @property
    def salary(self): return self.hours * self.rate

    @property
    def kpi(self): return self.sales * KPI_PERCENT / 100

    @property
    def total(self): return self.salary + self.kpi - self.option - self.ofzp

    @property
    def payout(self) -> int:
        return int(math.ceil(round(self.total, 2) / ROUND_TO) * ROUND_TO)


@dataclass
class Person:
    name: str
    entries: list = field(default_factory=list)
    period: str = ""

    @property
    def hours(self): return sum(e.hours for e in self.entries)

    @property
    def payout(self): return sum(e.payout for e in self.entries)


class ParseError(Exception):
    pass


def to_num(s) -> float:
    if s is None or s == "":
        return 0.0
    if isinstance(s, (int, float)):
        return float(s)
    t = str(s).replace(" ", "").replace(" ", "").replace(" ", "").replace("₽", "")
    t = t.replace("р", "").replace(",", ".").replace("−", "-")
    if not t:
        return 0.0
    return float(t)


PERIOD_RE = re.compile(r"(\d{1,2}\.\d{1,2}(?:\.\d{2,4})?)\s*[-–—]\s*(\d{1,2}\.\d{1,2}(?:\.\d{2,4})?)")


def find_period(text):
    m = PERIOD_RE.search(str(text or ""))
    return f"{m.group(1)} – {m.group(2)}" if m else None


def group(records):
    """[(name, Entry)] -> [Person], одна карточка на человека."""
    people = OrderedDict()
    for name, e in records:
        key = norm_name(name)
        people.setdefault(key, Person(" ".join(str(name).split())))
        people[key].entries.append(e)
    return list(people.values())


# ---------- ручной ввод ----------
_LABELS = {"оф": "ofzp", "офзп": "ofzp", "официалка": "ofzp", "опц": "option", "опция": "option"}


def parse_line(line: str):
    """«Стас С 55 66300», можно добавить «оф 13593» и/или «опц 500»."""
    extra = {}
    if re.search(r"[;\t|]", line):
        f = [x.strip() for x in re.split(r"[;\t|]", line)]
        f += [""] * (6 - len(f))
        name, pt, h, s, opt, of = f[:6]
        code = find_point(pt)
        if not code:
            raise ParseError(f"не знаю точку «{pt}»")
        return name, Entry(code, to_num(h), to_num(s), option=to_num(opt), ofzp=to_num(of))

    tok = line.split()
    rest = []
    i = 0
    while i < len(tok):
        lab = _norm(tok[i]).rstrip(":")
        if lab in _LABELS and i + 1 < len(tok):
            extra[_LABELS[lab]] = to_num(tok[i + 1])
            i += 2
            continue
        rest.append(tok[i])
        i += 1
    if len(rest) < 4:
        raise ParseError("нужно: имя точка часы продажи")
    try:
        h, s = to_num(rest[-2]), to_num(rest[-1])
    except ValueError:
        raise ParseError("после точки должны идти числа: часы и продажи")
    code = find_point(rest[-3])
    if not code:
        raise ParseError(f"не знаю точку «{rest[-3]}». Точки: " + ", ".join(POINTS))
    return " ".join(rest[:-3]), Entry(code, h, s, **extra)


def parse_text(text: str):
    records, errors, period = [], [], None
    for i, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if not line:
            continue
        p = find_period(line)
        if p and not re.sub(r"(?i)период|:", "", PERIOD_RE.sub("", line)).strip():
            period = p
            continue
        try:
            records.append(parse_line(line))
        except (ParseError, ValueError) as e:
            errors.append(f"Строка {i} «{line}»: {e}")
    return period, group(records), errors, []


# ---------- таблица ----------
def _col_key(h: str):
    h = re.sub(r"\s+", "", str(h).lower())
    if not h:
        return None
    if "период" in h: return "period"
    if any(k in h for k in ("сотруд", "имя", "фио")): return "name"
    if "точк" in h: return "point"
    if "час" in h: return "hours"
    if "ставк" in h: return "rate"
    if "продаж" in h: return "sales"
    if "опц" in h: return "option"
    if h.startswith("оф"): return "ofzp"
    if "выплат" in h: return "payout"
    return None


def _read_sheets(filename: str, data: bytes):
    """[(название листа, строки)]. Из xlsx берём листы «Зарплата …», иначе все."""
    fn = filename.lower()
    if fn.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheets = [ws for ws in wb.worksheets if "зарплат" in ws.title.lower()] or wb.worksheets
        return [(ws.title, [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)])
                for ws in sheets]
    if fn.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:3000], delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
        return [("", [row for row in csv.reader(io.StringIO(text), dialect)])]
    raise ParseError("пришли .xlsx или .csv")


def _parse_grid(title, grid, want_period, errors):
    """Записи одного листа за нужный (или последний) период."""
    header_idx, cols = None, {}
    for i, row in enumerate(grid[:20]):
        found = {}
        for j, c in enumerate(row):
            k = _col_key(c)
            if k and k not in found:
                found[k] = j
        if {"name", "point", "hours", "sales"} <= found.keys():
            header_idx, cols = i, found
            break
    if header_idx is None:
        return None, []

    where = f"{title}, " if title else ""
    by_period = OrderedDict()
    for i, row in enumerate(grid[header_idx + 1:], header_idx + 2):
        get = lambda k: (row[cols[k]] if k in cols and cols[k] < len(row) else "")
        name = str(get("name")).strip()
        if not name:
            continue
        per = find_period(get("period")) or ""
        try:
            code = find_point(get("point"))
            if not code:
                raise ParseError(f"не знаю точку \u00ab{get('point')}\u00bb")
            rate = to_num(get("rate")) or RATE
            exp = get("payout")
            e = Entry(code, to_num(get("hours")), to_num(get("sales")), rate=rate,
                      option=to_num(get("option")), ofzp=to_num(get("ofzp")),
                      expected=to_num(exp) if str(exp).strip() != "" else None)
            by_period.setdefault(per, []).append((name, e))
        except (ParseError, ValueError) as ex:
            errors.append(f"{where}строка {i} ({name}): {ex}")
    if not by_period:
        return None, []
    if want_period:
        if want_period not in by_period:
            errors.append(f"{title or 'Таблица'}: периода {want_period} нет, лист пропущен")
            return None, []
        return want_period, by_period[want_period]
    key = list(by_period)[-1]
    return key, by_period[key]


def parse_table(filename: str, data: bytes, want_period: str = None):
    errors, people, periods, found_any = [], [], [], False
    for title, grid in _read_sheets(filename, data):
        key, records = _parse_grid(title, grid, want_period, errors)
        if key is None and not records:
            continue
        found_any = True
        periods.append(key)
        for p in group(records):
            p.period = key
            people.append(p)
    if not found_any and not errors:
        raise ParseError("не нашёл заголовки. Нужны колонки: сотрудник, точка, часы, продажи")

    mismatches = []
    for p in people:
        for e in p.entries:
            if e.expected is not None and abs(e.expected - e.payout) >= 1:
                mismatches.append(f"{p.name} ({e.point}): в таблице {e.expected:g}, у меня {e.payout}")
    uniq = [x for x in dict.fromkeys(periods) if x]
    if len(uniq) > 1:
        errors.append("На листах разные последние периоды: " + ", ".join(uniq)
                      + ". Если нужен один \u2014 укажи его в подписи к файлу.")
    return (" / ".join(uniq) or None), people, errors, mismatches


def make_template() -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Выписка"
    ws.append(["период", "сотрудник", "точка", "часы", "ставка", "зп", "продажи", "кпи",
               "опция", "оф зп", "итог", "к выплате"])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4A5646")
    ws.append(["1.09-15.09", "Даша", "С", 28, 320, None, 0])
    ws.append(["1.09-15.09", "Василина", "В", 70, 320, None, 42120, None, None, 13593])
    ws.append(["1.09-15.09", "Оля", "П", 60, 320, None, 41700])
    ws.append(["1.09-15.09", "Оля", "В", 20, 320, None, 10000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
